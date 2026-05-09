import random
import os
from datetime import date

import openpyxl
from openpyxl.styles import Font
from google.cloud import storage

# ─────────────────────────────────────────────────────────────────────────────
# 1) BASIC CONFIG
# ─────────────────────────────────────────────────────────────────────────────

TODAY = date.today().strftime("%Y%m%d")
BUCKET_NAME = "mobile_brands"
LANDING_FOLDER = "landing"  # Added this to target the subfolder

TMP_DIR = "/tmp/distributor_files"
os.makedirs(TMP_DIR, exist_ok=True)

storage_client = storage.Client()
bucket = storage_client.bucket(BUCKET_NAME)

# ─────────────────────────────────────────────────────────────────────────────
# 2) BASE CATALOGUE (INDIA / INR)
# ─────────────────────────────────────────────────────────────────────────────

CATALOGUE_INR = {
    "Samsung": [
        ("Galaxy A56 5G", 38990, "1029384756102938"),
        ("Galaxy M17 5G", 13965, "1928374650192837"),
        ("Galaxy S26 Ultra", 121870, "1938475620193847"),
        ("Galaxy F07", 9499, "2039485710293847"),
        ("Galaxy S26", 76500, "2948571029384756"),
        ("Galaxy F34 5G", 21999, "3847562910384756"),
        ("Galaxy S24 Ultra 5G", 134999, "4729104857291038"),
        ("Galaxy A36 5G", 35208, "4857291038475629"),
        ("Galaxy M35 5G", 17499, "5729103847562910"),
        ("Galaxy S22 5G", 85999, "5829103847562910"),
        ("Galaxy M56 5G", 23499, "7584930218475629"),
        ("Galaxy Z Fold 5", 159999, "8392048175930284"),
        ("Galaxy A07", 10300, "8475629103847562"),
        ("Galaxy S24 FE", 39999, "9283746501928374"),
        ("Galaxy A06", 7999, "9384756201938475"),
    ],
    "Vivo": [
        ("X Fold3 Pro", 159999, "3829104857291030"),
        ("X100 Pro", 89999, "5729104857291031"),
        ("X80 Pro", 79999, "1938475620193841"),
        ("X90 Pro", 74999, "5829103847562911"),
        ("X100", 63999, "2948571029384751"),
        ("V30 Pro", 41999, "9283746501928371"),
        ("V30", 33999, "1029384756102931"),
        ("V29 5G", 32999, "4857291038475621"),
        ("T2 Pro 5G", 23999, "7584930218475621"),
        ("Y200 5G", 21999, "3847562910384751"),
        ("Y56 5G", 16999, "5729103847562911"),
        ("Y28 5G", 13999, "1928374650192831"),
        ("T2x 5G", 12999, "8475629103847561"),
        ("Y16", 10499, "2039485710293841"),
        ("Y02t", 8999, "9384756201938471"),
    ],
    "Oppo": [
        ("Find N3 Fold", 149999, "8392048175930281"),
        ("Find X7 Ultra", 99999, "4729104857291031"),
        ("Find N2 Flip", 89999, "1938475620193841"),
        ("Find X6 Pro", 84999, "5829103847562911"),
        ("Reno 10 Pro+ 5G", 54999, "2948571029384751"),
        ("Reno 11 Pro 5G", 39999, "9283746501928371"),
        ("Reno 10 5G", 32999, "1029384756102931"),
        ("Reno 11 5G", 29999, "4857291038475621"),
        ("F23 5G", 24999, "7584930218475621"),
        ("F25 Pro 5G", 23999, "3847562910384751"),
        ("A79 5G", 19999, "5729103847562911"),
        ("A59 5G", 14999, "1928374650192831"),
        ("A58", 13999, "8475629103847561"),
        ("A38", 12999, "2039485710293841"),
        ("A18", 9999, "9384756201938471"),
    ],
    "OnePlus": [
        ("OnePlus Open", 139999, "8392048175930280"),
        ("OnePlus 10 Pro", 66999, "4729104857291030"),
        ("OnePlus 12", 64999, "1938475620193840"),
        ("OnePlus 9 Pro", 64999, "5829103847562910"),
        ("OnePlus 11 5G", 56999, "2948571029384750"),
        ("OnePlus 12R", 39999, "9283746501928370"),
        ("OnePlus 11R", 35999, "1029384756102930"),
        ("OnePlus Nord 3", 33999, "4857291038475620"),
        ("OnePlus Nord 4", 29999, "7584930218475620"),
        ("OnePlus Nord CE 4", 24999, "3847562910384750"),
    ],
    "Apple": [
        ("iPhone 16 Pro Max", 159900, "8392048175930283"),
        ("iPhone 16 Pro", 134900, "4729104857291033"),
        ("iPhone 16", 90000, "1938475620193843"),
        ("iPhone 15 Pro Max", 139900, "5829103847562913"),
        ("iPhone 15 Pro", 114900, "2948571029384753"),
        ("iPhone 15 Plus", 89900, "3847562910384753"),
        ("iPhone 15", 79900, "9283746501928373"),
        ("iPhone 14 Pro Max", 129900, "1029384756102933"),
        ("iPhone 14 Pro", 99000, "4857291038475623"),
        ("iPhone 14", 60000, "7584930218475623"),
    ],
}

# ─────────────────────────────────────────────────────────────────────────────
# 3) COUNTRY CONFIG
# ─────────────────────────────────────────────────────────────────────────────

COUNTRIES = {
    "India": {"prefix": "IND", "currency": "INR", "fx": 1.0, "factor": 1.00, "gcs_folder": "INDIA"},
    "China": {"prefix": "CHN", "currency": "CNY", "fx": 0.086, "factor": 1.02, "gcs_folder": "CHINA"},
    "America": {"prefix": "USA", "currency": "USD", "fx": 0.012, "factor": 1.05, "gcs_folder": "AMERICA"},
    "South Korea": {"prefix": "KOR", "currency": "KRW", "fx": 16.2, "factor": 1.04, "gcs_folder": "SOUTH KOREA"},
    "Australia": {"prefix": "AUS", "currency": "AUD", "fx": 0.018, "factor": 1.03, "gcs_folder": "AUSTRALIA"},
    "South Africa": {"prefix": "ZAF", "currency": "ZAR", "fx": 0.23, "factor": 1.06, "gcs_folder": "SOUTH AFRICA"},
}

NUM_DISTRIBUTORS = 15
NUM_RETAILERS = 10
BRAND_ORDER = ["Samsung", "Vivo", "Oppo", "OnePlus", "Apple"]

HEADERS = [
    "Brand", "Model", "Price", "Distributor_code",
    "Retailer_code", "Store_code", "EAN_code",
    "Currency", "Stock_units", "Sale_units"
]

# ─────────────────────────────────────────────────────────────────────────────
# 4) HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def local_price(price_inr, fx, factor):
    return max(1, int(round(price_inr * fx * factor)))

def rand_stock(price, scale):
    if price >= 100000:
        s_min, s_max = 3, 20
    elif price >= 50000:
        s_min, s_max = 5, 35
    elif price >= 25000:
        s_min, s_max = 10, 60
    elif price >= 15000:
        s_min, s_max = 20, 90
    else:
        s_min, s_max = 40, 180

    stock = random.randint(int(s_min * scale), int(s_max * scale))
    sales = random.randint(1, stock)
    return stock, sales

# ─────────────────────────────────────────────────────────────────────────────
# 5) GENERATION + UPLOAD
# ─────────────────────────────────────────────────────────────────────────────

for country, cfg in COUNTRIES.items():
    print(f"\n=== {country} ===")

    for d in range(1, NUM_DISTRIBUTORS + 1):
        dist_code = f"{cfg['prefix']}DST{d:02d}"
        scale = 0.85 + (d - 1) * 0.03

        hierarchy = {}
        for r in range(1, NUM_RETAILERS + 1):
            ret = f"{dist_code}RET{r:02d}"
            hierarchy[ret] = [
                f"{ret}ST{s:02d}" for s in range(1, random.randint(5, 10) + 1)
            ]

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for brand in BRAND_ORDER:
            ws = wb.create_sheet(brand)
            for c, h in enumerate(HEADERS, 1):
                ws.cell(row=1, column=c, value=h).font = Font(bold=True)

            row = 2
            for ret, stores in hierarchy.items():
                for st in stores:
                    for model, price_inr, ean in CATALOGUE_INR[brand]:
                        price = local_price(price_inr, cfg["fx"], cfg["factor"])
                        su, sa = rand_stock(price, scale)
                        values = [
                            brand, model, price, dist_code,
                            ret, st, ean, cfg["currency"], su, sa
                        ]
                        for c, v in enumerate(values, 1):
                            ws.cell(row=row, column=c, value=v)
                        row += 1

        file_name = f"{dist_code}_{TODAY}.xlsx"
        local_path = os.path.join(TMP_DIR, file_name)
        wb.save(local_path)

        # ✅ UPDATED: Files now go into landing/<COUNTRY>/filename
        gcs_path = f"{LANDING_FOLDER}/{cfg['gcs_folder']}/{file_name}"
        bucket.blob(gcs_path).upload_from_filename(local_path)

        print(f"Uploaded → gs://{BUCKET_NAME}/{gcs_path}")

print("\n✅ ALL FILES SUCCESSFULLY UPLOADED TO LANDING BUCKET")